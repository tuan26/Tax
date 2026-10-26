"""Dựng lại tools/pilot/baseline.xlsx: python tools/pilot/build_baseline.py (cần openpyxl)."""
from datetime import date, time
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

F = "Arial"
BASE = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
HEAD = Font(name=F, size=10, bold=True, color="FFFFFF")
TITLE = Font(name=F, size=14, bold=True)
EXAMPLE = Font(name=F, size=10, italic=True, color="666666")
NOTE = Font(name=F, size=9, color="5D6964")
INPUT = PatternFill("solid", fgColor="FFF2CC")
HEADFILL = PatternFill("solid", fgColor="1F6B57")
SUMFILL = PatternFill("solid", fgColor="E1EFE9")
THIN = Border(bottom=Side(style="thin", color="D9DFDA"))
N = 300

STAGES = ["1. Nhận và tải chứng từ", "2. Nhập liệu", "3. Đối chiếu, kiểm tra", "4. Hỏi lại khách",
          "5. Chốt, nhập phần mềm kế toán", "6. Khác"]

wb = Workbook()


def sheet(title, headers, widths, example, input_cols, formulas=None, note=None):
    ws = wb.create_sheet(title)
    ws["A1"] = title
    ws["A1"].font = TITLE
    if note:
        ws["A2"] = note
        ws["A2"].font = NOTE
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(3, i, h)
        c.font, c.fill = HEAD, HEADFILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[3].height = 32
    for r in range(4, 4 + N):
        for i in range(1, len(headers) + 1):
            c = ws.cell(r, i)
            c.font = BASE
            c.border = THIN
            if i in input_cols:
                c.fill = INPUT
        for col, f in (formulas or {}).items():
            ws.cell(r, col, f.format(r=r))
    for i, v in enumerate(example, 1):
        if v is not None and i in input_cols:
            ws.cell(4, i, v).font = EXAMPLE
    ws.freeze_panes = "A4"
    return ws


def validation(ws, col, options_ref):
    dv = DataValidation(type="list", formula1=options_ref, allow_blank=True)
    dv.error = "Chọn trong danh sách"
    ws.add_data_validation(dv)
    dv.add(f"{col}4:{col}{3 + N}")


lists = wb.create_sheet("Danh mục")
cats = {
    "A": ("Khâu", STAGES),
    "B": ("Phương pháp thuế", ["Theo % doanh thu", "Theo thu nhập", "Chưa rõ"]),
    "C": ("Kênh Zalo", ["Chat 1-1", "Nhóm Zalo", "Cả hai", "Khác"]),
    "D": ("Có/Không", ["Có", "Không", "Chưa rõ"]),
    "E": ("Loại câu khó", ["Thứ tự (cái đầu, cái thứ 2)", "Tham chiếu (cái trên)", "Rút lại (gửi nhầm, không lấy)",
                          "Đính chính số", "Gửi bù, gửi muộn", "Nhầm hộ", "Ngày tương đối (hôm qua, tối qua)",
                          "Nhiều ảnh một giao dịch", "Gửi lại ảnh cũ", "Khác"]),
    "F": ("Phần mềm kế toán", ["MISA SME", "MISA ASP", "MISA AMIS", "Excel", "Khác"]),
}
for col, (name, items) in cats.items():
    lists[f"{col}1"] = name
    lists[f"{col}1"].font = BOLD
    for i, v in enumerate(items, 2):
        lists[f"{col}{i}"] = v
        lists[f"{col}{i}"].font = BASE
lists.sheet_state = "hidden"


def ref(col, n):
    return f"='Danh mục'!${col}$2:${col}${n + 1}"


log = sheet("Ghi thời gian",
            ["Ngày", "Kế toán", "Hộ", "Khâu", "Bắt đầu (giờ:phút)", "Kết thúc (giờ:phút)", "Số phút",
             "Số mục xử lý (ảnh, tin, giao dịch)", "Ghi chú"],
            [12, 16, 22, 30, 14, 14, 10, 18, 40],
            [date(2026, 10, 27), "Chị Hà", "Quán Minh Anh", STAGES[0], time(8, 30), time(8, 52), None, 37,
             "Tải 37 ảnh từ Zalo PC, đặt tên theo ngày (ví dụ, xóa dòng này)"],
            {1, 2, 3, 4, 5, 6, 8, 9},
            {7: '=IF(OR(E{r}="",F{r}=""),"",ROUND(MOD(F{r}-E{r},1)*1440,0))'},
            note="Mỗi lần làm việc với một hộ ghi một dòng. Ô vàng là ô cần điền. Số phút tự tính, kể cả khi làm qua nửa đêm.")
for r in range(4, 4 + N):
    log.cell(r, 1).number_format = "dd/mm/yyyy"
    log.cell(r, 5).number_format = "hh:mm"
    log.cell(r, 6).number_format = "hh:mm"
validation(log, "D", ref("A", len(STAGES)))

prof = sheet("Hồ sơ hộ",
             ["Hộ", "Kế toán", "Ngành", "Phương pháp thuế", "Kênh gửi chứng từ", "Số người gửi", "Tin nhắn/tháng",
              "Ảnh/tháng", "Có phần mềm bán hàng (POS)?", "Có hóa đơn điện tử?", "Phần mềm kế toán đích",
              "Phí dịch vụ/tháng (đ)", "Ghi chú"],
             [22, 16, 18, 18, 16, 10, 12, 10, 14, 14, 16, 16, 36],
             ["Quán Minh Anh", "Chị Hà", "Ăn uống", "Theo % doanh thu", "Nhóm Zalo", 3, 210, 83, "Không", "Không",
              "MISA SME", 800000, "Chủ hộ, vợ và thu ngân cùng gửi (ví dụ)"],
             set(range(1, 14)),
             note="Một dòng cho mỗi hộ. Tên hộ ở đây được dùng ở sheet Tổng hợp.")
for col, (lst, n) in {"D": ("B", 3), "E": ("C", 4), "I": ("D", 3), "J": ("D", 3), "K": ("F", 5)}.items():
    validation(prof, col, ref(lst, n))
for r in range(4, 4 + N):
    prof.cell(r, 12).number_format = "#,##0"

summ = wb.create_sheet("Tổng hợp")
summ["A1"] = "Tổng hợp baseline"
summ["A1"].font = TITLE
summ["A2"] = ("Tự tính từ sheet Ghi thời gian. Chỉ điền cột 'Số ngày đã đo' (ô vàng): số ngày làm việc đã ghi đủ "
              "cho hộ đó, để quy ra phút mỗi tháng.")
summ["A2"].font = NOTE
heads = ["Hộ"] + [s.split(". ", 1)[1] for s in STAGES] + ["Tổng số phút đã đo", "Số ngày đã đo",
                                                           "Ước tính phút/hộ/tháng"]
for i, h in enumerate(heads, 1):
    c = summ.cell(3, i, h)
    c.font, c.fill = HEAD, HEADFILL
    c.alignment = Alignment(wrap_text=True, vertical="center")
    summ.column_dimensions[get_column_letter(i)].width = 22 if i == 1 else 14
summ.row_dimensions[3].height = 44
summ["L3"] = "Số ngày làm việc/tháng"
summ["L3"].font = BOLD
summ["M3"] = 22
summ["M3"].font = Font(name=F, size=10, color="0000FF")
summ["M3"].fill = INPUT
summ["L4"] = "Giả định để quy đổi; sửa nếu văn phòng làm 26 ngày"
summ["L4"].font = NOTE
summ.column_dimensions["L"].width = 24
ROWS = 40
for r in range(4, 4 + ROWS):
    summ.cell(r, 1, f"=IF('Hồ sơ hộ'!A{r}=\"\",\"\",'Hồ sơ hộ'!A{r})").font = Font(name=F, size=10, color="008000")
    for k, stage in enumerate(STAGES):
        summ.cell(r, 2 + k, f"=IF($A{r}=\"\",\"\",SUMIFS('Ghi thời gian'!$G$4:$G${3 + N},"
                            f"'Ghi thời gian'!$C$4:$C${3 + N},$A{r},'Ghi thời gian'!$D$4:$D${3 + N},\"{stage}\"))").font = BASE
    summ.cell(r, 8, f"=IF($A{r}=\"\",\"\",SUM(B{r}:G{r}))").font = BOLD
    c = summ.cell(r, 9)
    c.fill, c.font = INPUT, BASE
    summ.cell(r, 10, f"=IF(OR($A{r}=\"\",I{r}=\"\",I{r}=0),\"\",ROUND(H{r}/I{r}*$M$3,0))").font = BOLD
    for col in range(1, 11):
        summ.cell(r, col).border = THIN
    summ.cell(r, 10).fill = SUMFILL
summ.cell(4, 9, 1).font = EXAMPLE
tot = 4 + ROWS
summ.cell(tot, 1, "Tổng các hộ").font = BOLD
for col in range(2, 9):
    L = get_column_letter(col)
    summ.cell(tot, col, f"=SUM({L}4:{L}{tot - 1})").font = BOLD
summ.freeze_panes = "B4"

hard = sheet("Câu khó",
             ["Ngày", "Hộ", "Nguyên văn tin nhắn", "Ngữ cảnh (tin ngay trước và sau)", "Nghĩa đúng (kế toán giải thích)",
              "Loại", "Kế toán mất bao lâu để hiểu (phút)", "Đã thành ca kiểm thử?"],
             [12, 20, 34, 40, 40, 26, 14, 14],
             [date(2026, 10, 27), "Quán Minh Anh", "cái đầu 2tr5", "Trước đó gửi 2 ảnh hóa đơn liền nhau",
              "2.500.000 là tổng của ảnh thứ nhất (ví dụ)", "Thứ tự (cái đầu, cái thứ 2)", 2, "Không"],
             {1, 2, 3, 4, 5, 6, 7, 8},
             note="Mỗi câu khó hiểu ghi một dòng. Đây là dữ liệu quan trọng nhất cho engine gom tin.")
for r in range(4, 4 + N):
    hard.cell(r, 1).number_format = "dd/mm/yyyy"
validation(hard, "F", ref("E", 10))
validation(hard, "H", ref("D", 3))

miss = sheet("Bỏ sót",
             ["Ngày phát hiện", "Hộ", "Mô tả chứng từ thiếu hoặc sai", "Lẽ ra phát hiện từ ngày", "Phát hiện bằng cách nào",
              "Ảnh hưởng (số tiền, phải sửa sổ, nộp bổ sung…)"],
             [14, 20, 40, 16, 30, 36],
             [date(2026, 10, 27), "Quán Minh Anh", "Chuyển khoản 5.000.000đ ngày 03/10 không có hóa đơn",
              date(2026, 10, 3), "Đối chiếu sao kê cuối tháng", "Phải hỏi lại khách, mất 2 ngày (ví dụ)"],
             {1, 2, 3, 4, 5, 6},
             note="Chứng từ thiếu hoặc sai phát hiện muộn. Dùng để so: trong pilot, app có bỏ sót nhiều hơn cách làm cũ không.")
for r in range(4, 4 + N):
    miss.cell(r, 1).number_format = "dd/mm/yyyy"
    miss.cell(r, 4).number_format = "dd/mm/yyyy"

guide = wb["Sheet"]
guide.title = "Hướng dẫn"
lines = [
    ("Phiếu baseline trước pilot", TITLE),
    ("Mục đích: đo cách làm hiện tại để sau pilot so được trước và sau, trên cùng hộ.", BASE),
    ("", BASE),
    ("Cách dùng", BOLD),
    ("Ô nền vàng là ô cần điền. Ô khác có công thức, không sửa.", BASE),
    ("Dòng chữ nghiêng màu xám ở mỗi sheet là ví dụ cách ghi. Xóa nội dung dòng đó trước khi dùng.", BASE),
    ("1. Hồ sơ hộ: một dòng cho mỗi hộ. Tên hộ ở đây tự xuất hiện ở sheet Tổng hợp.", BASE),
    ("2. Ghi thời gian: khi ngồi cạnh kế toán, mỗi lần họ làm việc với một hộ ghi một dòng theo khâu.", BASE),
    ("3. Tổng hợp: tự tính. Chỉ điền 'Số ngày đã đo' cho mỗi hộ để quy ra phút mỗi tháng.", BASE),
    ("4. Câu khó: mọi tin nhắn kế toán phải dừng lại để hiểu. Mỗi dòng sẽ thành một ca kiểm thử của engine.", BASE),
    ("5. Bỏ sót: chứng từ thiếu hoặc sai phát hiện muộn.", BASE),
    ("", BASE),
    ("Các khâu", BOLD),
] + [(s, BASE) for s in STAGES] + [
    ("", BASE),
    ("Các khâu khớp với chỉ số app đo trong pilot (python -m app.metrics), nên so được trực tiếp.", BASE),
    ("Tài liệu kèm theo trong repo: docs/pilot/README.md, phong-van-ke-toan.md, chep-chat-zalo.md.", BASE),
]
for i, (text, font) in enumerate(lines, 1):
    guide.cell(i, 1, text).font = font
guide.column_dimensions["A"].width = 110

wb._sheets = [wb[n] for n in ["Hướng dẫn", "Ghi thời gian", "Tổng hợp", "Hồ sơ hộ", "Câu khó", "Bỏ sót", "Danh mục"]]
wb.active = 0
# File không có giá trị tính sẵn: buộc Excel, LibreOffice, Google Sheets tính lại toàn bộ khi mở.
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(__import__("pathlib").Path(__file__).with_name("baseline.xlsx"))
print("saved")
