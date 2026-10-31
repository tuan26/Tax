"""Xuất Excel: tổng chính thức chỉ cộng giao dịch đã xác nhận; dữ liệu khách nhập không thành công thức."""

import io

import psycopg
from openpyxl import load_workbook

from tests.conftest import login
from tests.test_integration import new_customer, png


def _xlsx(client, h, cid, **params):
    r = client.get(f"/customers/{cid}/export.xlsx", headers=h, params=params)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    return load_workbook(io.BytesIO(r.content))


def _summary(ws):
    return {ws.cell(i, 1).value: ws.cell(i, 4).value for i in range(4, 10)}


def _lines(ws):
    head = [c.value for c in ws[11]]
    return [dict(zip(head, (c.value for c in row))) for row in ws.iter_rows(min_row=12) if row[0].value is not None]


def test_export_counts_only_confirmed(client, pg, tenants):
    h = login(client, pg, "A", tenants)
    cid = new_customer(client, h, "Quán xuất Excel")
    r = client.post(f"/customers/{cid}/imports", headers=h, data={
        "business_date": "2026-10-05", "chat_text": "Doanh thu 5tr\ntiền gas"},
        files=[("files", ("hd.png", io.BytesIO(png(250)), "image/png"))])
    assert r.status_code == 201, r.text
    recs = client.get(f"/customers/{cid}/records", headers=h).json()
    income = next(x for x in recs if x["evidence_kind"] == "self_declared")
    image = next(x for x in recs if x["evidence_kind"] == "document")
    assert not income["confirmed"] and not income["review_required"], "doanh thu tự khai: tự động, chưa xác nhận"
    assert image["review_required"], "ảnh không có OCR: chờ duyệt"

    wb = _xlsx(client, h, cid, date_from="2026-10-01", date_to="2026-10-31")
    ws = wb["So theo ngay"]
    s = _summary(ws)
    assert s["Tổng thu đã xác nhận"] == 0, "chưa xác nhận thì không vào tổng chính thức"
    assert s["Thu tự động, chưa xác nhận (không tính vào tổng trên)"] == 5_000_000
    statuses = {str(x["Mã giao dịch"]): x["Trạng thái"] for x in _lines(ws)}
    assert statuses[str(income["id"])] == "Tự động, chưa xác nhận"
    assert statuses[str(image["id"])] == "Chờ duyệt"

    # Kế toán xác nhận doanh thu; nội dung có dạng công thức Excel.
    r = client.post(f"/records/{income['id']}/confirm", headers=h,
                    json={"fields": {"description": '=HYPERLINK("http://x.vn","bấm")'}})
    assert r.status_code == 200, r.text
    ws = _xlsx(client, h, cid, date_from="2026-10-01", date_to="2026-10-31")["So theo ngay"]
    s = _summary(ws)
    assert s["Tổng thu đã xác nhận"] == 5_000_000
    assert s["Thu tự động, chưa xác nhận (không tính vào tổng trên)"] == 0
    line = next(x for x in _lines(ws) if str(x["Mã giao dịch"]) == str(income["id"]))
    assert line["Tính vào tổng"] == "Có" and line["Số tiền"] == 5_000_000
    cell = next(c for row in ws.iter_rows(min_row=12) for c in row if c.value == '=HYPERLINK("http://x.vn","bấm")')
    assert cell.data_type == "s", "chuỗi do người dùng nhập không được thành công thức"

    # Ngoài kỳ thì không có trong file.
    ws = _xlsx(client, h, cid, date_from="2026-11-01", date_to="2026-11-30")["So theo ngay"]
    assert _lines(ws) == []

    # Mỗi lần xuất có audit: ai, hộ nào, kỳ nào.
    tid, uid = tenants["A"]
    with psycopg.connect(pg["admin"]) as conn:
        rows = conn.execute("SELECT actor_id, new_row FROM audit_event WHERE action = 'EXPORT_XLSX' AND row_id = %s"
                            " ORDER BY id", (str(cid),)).fetchall()
    assert len(rows) == 3 and all(a == uid for a, _ in rows)
    assert rows[0][1]["date_from"] == "2026-10-01"


def test_export_lists_open_findings_and_other_tenant_cannot_export(client, pg, tenants):
    ha = login(client, pg, "A", tenants)
    cid = new_customer(client, ha, "Hộ có việc cần xử lý")
    client.post(f"/customers/{cid}/imports", headers=ha, data={"business_date": "2026-10-06", "chat_text": "Doanh thu 2tr"})
    wb = _xlsx(client, ha, cid)
    assert [c.value for c in wb["Viec can xu ly"][1]][:2] == ["Mã rule", "Việc cần xử lý"]
    assert wb["Thong tin"]["B5"].value is not None

    hb = login(client, pg, "B", tenants)
    r = client.get(f"/customers/{cid}/export.xlsx", headers=hb)
    assert r.status_code == 404
    with psycopg.connect(pg["admin"]) as conn:
        n = conn.execute("SELECT count(*) FROM audit_event WHERE action = 'EXPORT_XLSX' AND row_id = %s"
                         " AND tenant_id = %s", (str(cid), tenants["B"][0])).fetchone()[0]
    assert n == 0
