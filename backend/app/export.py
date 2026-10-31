"""Xuất sổ theo ngày của một hộ ra Excel.

Tổng chính thức chỉ cộng giao dịch kế toán đã xác nhận. Giao dịch tự động (độ tin cậy cao, chưa xác nhận)
và giao dịch chờ duyệt vẫn có trong file, ghi rõ trạng thái, cộng ở dòng riêng và không vào tổng chính thức.
Khoản chuyển khoản đã ghép với chứng từ chỉ tính một lần (tính chứng từ), như màn hình sổ.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VN = ZoneInfo("Asia/Ho_Chi_Minh")
TYPE_LABEL = {"INCOME": "Thu", "EXPENSE": "Chi"}
KIND_LABEL = {"document": "Chứng từ", "bank_transfer_screenshot": "Chuyển khoản", "self_declared": "Tự khai"}
FINDING_STATUS = {"OPEN": "Đang mở", "REQUESTED": "Đã yêu cầu khách"}
MONEY = "#,##0"
HEAD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="E8EEEA")


def _day(v):
    try:
        return date.fromisoformat(v) if isinstance(v, str) else v
    except ValueError:
        return v


def status_of(row: dict) -> str:
    if row["confirmed"] and not row["review_required"]:
        return "Đã xác nhận"
    if row["review_required"]:
        return "Chờ duyệt"
    return "Tự động, chưa xác nhận"


def _text(ws, cell_ref, value):
    """Ghi chuỗi do người dùng nhập. Chuỗi bắt đầu bằng = không được thành công thức."""
    cell = ws[cell_ref]
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def _write_table(ws, start_row: int, headers: list[str], rows: list[list], money_cols=(), widths=()):
    for i, h in enumerate(headers, 1):
        c = ws.cell(start_row, i, h)
        c.font, c.fill = HEAD, HEAD_FILL
    for r, values in enumerate(rows, start_row + 1):
        for i, v in enumerate(values, 1):
            ref = f"{get_column_letter(i)}{r}"
            c = _text(ws, ref, v)
            if i in money_cols:
                c.number_format = MONEY
            elif isinstance(v, date):
                c.number_format = "dd/mm/yyyy"
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(start_row + 1, 1)


def build_workbook(customer: dict, ledger: list[dict], open_findings: list[dict], date_from, date_to,
                   exported_by: str) -> bytes:
    totals = {(t, s): 0 for t in TYPE_LABEL for s in ("official", "auto", "pending")}
    lines = []
    for n, r in enumerate(sorted(ledger, key=lambda x: (x["fields"].get("posting_date") or "", str(x["id"]))), 1):
        f = r["fields"]
        status = status_of(r)
        amount, ttype = f.get("amount"), f.get("transaction_type")
        if not r.get("counted", True):
            counted = "Không (đã ghép với chứng từ)"
        elif status == "Đã xác nhận":
            counted = "Có"
        else:
            counted = "Không (chưa xác nhận)"
        if amount is not None and ttype in TYPE_LABEL and r.get("counted", True):
            key = "official" if status == "Đã xác nhận" else ("pending" if status == "Chờ duyệt" else "auto")
            totals[(ttype, key)] += amount
        lines.append([n, _day(f.get("posting_date")), TYPE_LABEL.get(ttype, "Chưa rõ"), amount, f.get("description"),
                      _day(f.get("document_date")), KIND_LABEL.get(r.get("evidence_kind"), r.get("evidence_kind")), status,
                      counted, "; ".join(r.get("review_reasons") or []) or None, str(r["id"])])

    wb = Workbook()
    ws = wb.active
    ws.title = "So theo ngay"
    period = f"{date_from or '…'} đến {date_to or '…'}"
    _text(ws, "A1", f"Sổ theo ngày: {customer['name']}").font = Font(bold=True, size=13)
    _text(ws, "A2", f"Kỳ: {period}")
    summary = [
        ("Tổng thu đã xác nhận", totals[("INCOME", "official")]),
        ("Tổng chi đã xác nhận", totals[("EXPENSE", "official")]),
        ("Thu tự động, chưa xác nhận (không tính vào tổng trên)", totals[("INCOME", "auto")]),
        ("Chi tự động, chưa xác nhận (không tính vào tổng trên)", totals[("EXPENSE", "auto")]),
        ("Thu chờ duyệt (không tính)", totals[("INCOME", "pending")]),
        ("Chi chờ duyệt (không tính)", totals[("EXPENSE", "pending")]),
    ]
    for i, (label, value) in enumerate(summary, 4):
        ws.cell(i, 1, label).font = HEAD if i < 6 else Font()
        ws.cell(i, 4, value).number_format = MONEY
    _write_table(ws, 11, ["STT", "Ngày hạch toán", "Thu/Chi", "Số tiền", "Nội dung", "Ngày chứng từ", "Bằng chứng",
                          "Trạng thái", "Tính vào tổng", "Lý do chờ duyệt", "Mã giao dịch"], lines,
                 money_cols=(4,), widths=(6, 14, 9, 14, 40, 14, 14, 22, 26, 30, 38))

    wf = wb.create_sheet("Viec can xu ly")
    _write_table(wf, 1, ["Mã rule", "Việc cần xử lý", "Chi tiết", "Số tiền", "Trạng thái", "Mã giao dịch"],
                 [[f["rule_id"], f["title"], f["detail"], f["amount"], FINDING_STATUS.get(f["status"], f["status"]),
                   str(f["record_id"])] for f in open_findings],
                 money_cols=(4,), widths=(10, 34, 60, 14, 18, 38))

    wi = wb.create_sheet("Thong tin")
    info = [["Hộ", customer["name"]], ["Kỳ", period],
            ["Thời điểm xuất", datetime.now(VN).strftime("%Y-%m-%d %H:%M (giờ Việt Nam)")],
            ["Người xuất", exported_by],
            ["Số giao dịch", len(lines)],
            ["Ghi chú", "Tổng chính thức chỉ cộng giao dịch đã xác nhận. Khoản chuyển khoản đã ghép với chứng từ "
                        "chỉ tính một lần. File này là dữ liệu hỗ trợ kế toán, không phải tờ khai thuế."]]
    _write_table(wi, 1, ["Mục", "Giá trị"], info, widths=(18, 90))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
